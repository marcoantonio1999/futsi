from datetime import timedelta
from io import BytesIO
from unittest.mock import patch

import pytest
from django.utils import timezone
from openpyxl import load_workbook

from core.models import AuditLog, WhatsAppAutomationSettings, WhatsAppConversation, WhatsAppMessage, WhatsAppHumanResponseEvent
from core.tests.factories import make_site

pytestmark = [pytest.mark.api, pytest.mark.django_db]
BASE = "/api/whatsapp-conversations/"
A, B, C = [f"whatsapp:+52550000010{i}" for i in (1, 2, 3)]
FRANCO_ACADEMY = "meta:105039749242267"


def conversation(address, site=None):
    return WhatsAppConversation.objects.create(to_address=address, site=site,
        contact_phone="+525511110000", from_address="whatsapp:+525511110000", current_step="faq",
        last_message_at=timezone.now())


def test_all_sites_and_explicit_number_are_not_limited_to_env(auth_client, settings):
    client, _, _ = auth_client()
    north, south = make_site(), make_site()
    settings.META_WHATSAPP_DISPLAY_NUMBER = A.replace("whatsapp:", "")
    WhatsAppAutomationSettings.objects.create(business_address=A, site=north)
    WhatsAppAutomationSettings.objects.create(business_address=B, site=south)
    first, second, unknown = conversation(A, north), conversation(B, north), conversation(C)
    rows = client.get(BASE, {"scope": "all"}).json()
    assert {row["id"] for row in rows} == {first.id, second.id, unknown.id}
    row = next(row for row in rows if row["id"] == second.id)
    assert row["business_address"] == B and row["channel_site"] == south.id
    assert row["manual_send_available"] is False
    for query, ids in [({"site": south.id}, [second.id]), ({"scope": "all", "business_address": A}, [first.id]),
                       ({"scope": "all", "site": "unassigned"}, [unknown.id]),
                       ({"scope": "all", "site": south.id, "business_address": A}, [])]:
        assert [r["id"] for r in client.get(BASE, query).json()] == ids
    assert client.get(BASE, {"business_address": "bad"}).status_code == 400
    assert client.get(BASE, {"scope": "all", "site": "bad"}).status_code == 400


def test_channels_include_empty_sites_and_enforce_coordinator_scope(auth_client, api_client, settings):
    settings.META_WHATSAPP_DISPLAY_NUMBER = ""
    settings.META_WHATSAPP_PHONE_NUMBER_ID = ""
    north, south = make_site(), make_site()
    WhatsAppAutomationSettings.objects.create(business_address=A, site=north)
    WhatsAppAutomationSettings.objects.create(business_address=B, site=south)
    WhatsAppAutomationSettings.objects.create(business_address=C, site=south)
    # A legacy wrong default must not expose the new channel to the old coordinator.
    foreign = conversation(B, north)
    own = conversation(A)
    admin, _, _ = auth_client()
    assert {r["business_address"] for r in admin.get(BASE + "channels/").json()} == {A, B, C}
    coord, _, _ = auth_client(role="site_coordinator", primary_site=north)
    assert [r["id"] for r in coord.get(BASE, {"scope": "all"}).json()] == [own.id]
    assert [r["business_address"] for r in coord.get(BASE + "channels/").json()] == [A]
    assert coord.get(BASE + f"{foreign.id}/?scope=all").status_code == 404
    assert coord.patch(BASE + f"{foreign.id}/?scope=all", {"follow_up_required": True}, format="json").status_code == 404
    assert coord.get(BASE, {"scope": "all", "site": south.id}).json() == []
    no_site, _, _ = auth_client(role="site_coordinator")
    assert no_site.get(BASE + "channels/").json() == []
    assert api_client.get(BASE + "channels/").status_code in (401, 403)


def test_channels_include_configured_number_without_history_for_admin(auth_client, settings):
    settings.META_WHATSAPP_DISPLAY_NUMBER = A.replace("whatsapp:", "")
    settings.DUALHOOK_API_KEY = "local-key"
    settings.DUALHOOK_WABA_ID = "1234"
    settings.WHATSAPP_SERVICE_URL = ""
    settings.WHATSAPP_SERVICE_TOKEN = ""
    admin, _, _ = auth_client()

    assert admin.get(BASE + "channels/").json() == [{
        "business_address": A,
        "site": None,
        "site_name": "",
        "channel_label": "",
        "template_management_available": False,
    }]


def test_meta_channel_can_be_linked_and_filtered_inside_one_site(auth_client, settings):
    settings.META_WHATSAPP_DISPLAY_NUMBER = ""
    settings.META_WHATSAPP_PHONE_NUMBER_ID = ""
    settings.WHATSAPP_SERVICE_URL = ""
    settings.WHATSAPP_SERVICE_TOKEN = ""
    client, _, _ = auth_client()
    franco = make_site(name="Colegio Franco")
    WhatsAppAutomationSettings.objects.create(
        business_address=FRANCO_ACADEMY,
        channel_label="Franco Academia",
        site=franco,
    )
    chat = conversation(FRANCO_ACADEMY)
    veronica = conversation("meta:1100529023150528")
    veronica.context = {"kind": "veronica_manual"}
    veronica.save(update_fields=["context", "updated_at"])

    channels = client.get(BASE + "channels/").json()
    assert channels == [{
        "business_address": FRANCO_ACADEMY,
        "site": franco.id,
        "site_name": "Colegio Franco",
        "channel_label": "Franco Academia",
        "template_management_available": False,
    }]
    assert [row["id"] for row in client.get(BASE, {"scope": "all"}).json()] == [chat.id]
    rows = client.get(BASE, {
        "scope": "all",
        "site": franco.id,
        "business_address": FRANCO_ACADEMY,
    }).json()
    assert [row["id"] for row in rows] == [chat.id]
    assert rows[0]["channel_site"] == franco.id
    assert rows[0]["channel_label"] == "Franco Academia"


def test_weekly_statistics_use_the_same_site_and_channel_scope(auth_client):
    client, _, _ = auth_client()
    sites = [make_site(), make_site()]
    for address, site, duration in zip((A, B), sites, (60, 600)):
        WhatsAppAutomationSettings.objects.create(business_address=address, site=site)
        chat = conversation(address, site)
        message = WhatsAppMessage.objects.create(conversation=chat, direction="inbound", body="Horario?", contact_type="prospect")
        WhatsAppHumanResponseEvent.objects.create(conversation=chat, first_inbound_message=message,
            first_inbound_at=timezone.now(), responded_at=timezone.now() + timedelta(seconds=duration), response_seconds=duration)
    total = client.get(BASE + "weekly-stats/", {"scope": "all"}).json()
    assert total["summary"]["total"] == 2 and total["summary"]["average_response_seconds"] == 330
    for site, address, duration in zip(sites, (A, B), (60, 600)):
        row = client.get(BASE + "weekly-stats/", {"scope": "all", "site": site.id, "business_address": address}).json()
        assert row["summary"]["total"] == 1
        assert row["summary"]["average_response_seconds"] == duration
        assert row["classifications"]["prospect"] == 1
    assert client.get(BASE + "weekly-stats/", {"scope": "all", "business_address": C}).json()["summary"]["total"] == 0


def test_chat_export_contains_contact_summary_and_complete_message_detail(auth_client):
    client, _, user = auth_client()
    north, south = make_site(name="Norte"), make_site(name="Sur")
    WhatsAppAutomationSettings.objects.create(business_address=A, site=north, channel_label="Academia")
    WhatsAppAutomationSettings.objects.create(business_address=B, site=south, channel_label="Liga")
    academy = conversation(A, north)
    academy.context = {"contact_name": "=Nombre inseguro"}
    academy.save(update_fields=["context", "updated_at"])
    league = conversation(B, south)
    league.contact_phone = "+525522220000"
    league.context = {"contact_name": "Santiago Rivera"}
    league.save(update_fields=["contact_phone", "context", "updated_at"])
    WhatsAppMessage.objects.create(conversation=academy, direction="inbound", body="Quiero informes")
    WhatsAppMessage.objects.create(conversation=academy, direction="outbound", body="Con gusto")
    WhatsAppMessage.objects.create(conversation=league, direction="inbound", body="Quiero registrar mi equipo")

    response = client.get(BASE + "export/", {"scope": "all"})

    assert response.status_code == 200
    assert response["Content-Type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert "chats-whatsapp-futsi" in response["Content-Disposition"]
    workbook = load_workbook(BytesIO(response.content), data_only=False)
    assert workbook.sheetnames == ["Conversaciones", "Mensajes"]
    summary = workbook["Conversaciones"]
    headers = {cell.value: cell.column for cell in summary[1]}
    assert summary.max_row == 3
    academy_row = next(row for row in range(2, 4) if summary.cell(row, headers["Identificador de atención"]).value == A)
    assert summary.cell(academy_row, headers["Número de atención"]).value == "Academia"
    assert summary.cell(academy_row, headers["Nombre del contacto"]).value == "'=Nombre inseguro"
    assert summary.cell(academy_row, headers["Mensajes recibidos"]).value == 1
    assert summary.cell(academy_row, headers["Mensajes enviados"]).value == 1
    detail = workbook["Mensajes"]
    assert detail.max_row == 4
    assert {detail.cell(row, 6).value for row in range(2, 5)} == {"Recibido", "Enviado"}
    assert {detail.cell(row, 7).value for row in range(2, 5)} == {"Quiero informes", "Con gusto", "Quiero registrar mi equipo"}
    assert AuditLog.objects.filter(actor=user, action="whatsapp_chats_exported").exists()

    filtered = client.get(BASE + "export/", {"scope": "all", "business_address": B})
    filtered_workbook = load_workbook(BytesIO(filtered.content))
    assert filtered_workbook["Conversaciones"].max_row == 2
    assert filtered_workbook["Conversaciones"]["B2"].value == B


def test_chat_export_requires_an_admin_role(auth_client):
    site = make_site()
    coordinator, _, _ = auth_client(role="site_coordinator", primary_site=site)
    assert coordinator.get(BASE + "export/", {"scope": "all"}).status_code == 403


def test_manual_send_never_uses_another_sites_number(auth_client, settings):
    client, _, _ = auth_client()
    settings.META_WHATSAPP_DISPLAY_NUMBER = A.replace("whatsapp:", "")
    foreign = conversation(B)
    with patch("core.api.trials.send_text_for_channel") as send:
        response = client.post(BASE + f"{foreign.pk}/send-message/?scope=all", {"body": "Hola"}, format="json")
    assert response.status_code == 409
    send.assert_not_called()
    assert foreign.messages.count() == 0


def test_registered_service_channel_can_reply_from_its_own_number(auth_client, settings):
    client, _, _ = auth_client()
    settings.META_WHATSAPP_DISPLAY_NUMBER = A.replace("whatsapp:", "")
    settings.WHATSAPP_SERVICE_URL = "https://whatsapp-service.example"
    settings.WHATSAPP_SERVICE_TOKEN = "server-secret"
    franco = make_site(name="Colegio Franco")
    WhatsAppAutomationSettings.objects.create(
        business_address=FRANCO_ACADEMY,
        site=franco,
        channel_label="Franco Academia",
        bot_enabled=False,
    )
    chat = conversation(FRANCO_ACADEMY, franco)
    WhatsAppMessage.objects.create(
        conversation=chat,
        provider_sid="wamid.inbound-franco",
        direction="inbound",
        body="Hola, necesito información",
    )

    listed = client.get(BASE, {"scope": "all", "business_address": FRANCO_ACADEMY})
    assert listed.status_code == 200
    assert listed.json()[0]["manual_send_available"] is True

    with patch(
        "core.api.trials.send_text_for_channel",
        return_value="wamid.manual-franco",
    ) as send:
        response = client.post(
            BASE + f"{chat.pk}/send-message/?scope=all",
            {"body": "Hola, ¿en qué podemos ayudarte?"},
            format="json",
        )

    assert response.status_code == 201
    send.assert_called_once_with(
        address=FRANCO_ACADEMY,
        to_phone=chat.contact_phone,
        body="Hola, ¿en qué podemos ayudarte?",
    )
    assert response.json()["business_address"] == FRANCO_ACADEMY
    assert response.json()["human_takeover_active"] is True
    assert WhatsAppMessage.objects.filter(
        conversation=chat,
        provider_sid="wamid.manual-franco",
        direction="outbound",
    ).exists()
