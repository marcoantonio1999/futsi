from datetime import datetime, timedelta, timezone as dtz
from io import BytesIO
import pytest
from django.db import connection
from django.test.utils import CaptureQueriesContext
from openpyxl import load_workbook
from rest_framework.test import APIClient
from core.exam_models import RecruitmentExamAttempt
from core.tests.test_exam_profiles import attempt
from core.tests.test_recruitment_exams import create_exam

pytestmark=pytest.mark.django_db
URL='/api/recruitment-exams/export/'


def book(response):
    assert response.status_code==200,response.content
    assert response['Cache-Control']=='private, no-store'
    assert '.xlsx' in response['Content-Disposition']
    return load_workbook(BytesIO(response.content))


def test_all_pages_all_exams_and_read_only(auth_client):
    c,_,_=auth_client(role='collaborator',section_permissions=['veronica_only'])
    first=create_exam(c);second=create_exam(c)
    for i in range(65):attempt(first,i,0 if i==0 else 9.25)
    attempt(second,70,8)
    pending=attempt(first,71,None,submitted=False)
    # Even an expired, unsubmitted attempt must not be mutated by export.
    RecruitmentExamAttempt.objects.filter(pk=pending.pk).update(started_at=datetime(2020,1,1,tzinfo=dtz.utc))
    before=list(RecruitmentExamAttempt.objects.order_by('id').values())
    with CaptureQueriesContext(connection) as queries:
        w=book(c.get(URL))
    assert w.sheetnames==['Evaluaciones','Respuestas']
    assert w['Evaluaciones'].max_row==67
    assert len(queries)<6
    sql=' '.join(q['sql'].lower() for q in queries)
    assert not any(word in sql for word in ['update ','insert ','delete '])
    assert before==list(RecruitmentExamAttempt.objects.order_by('id').values())
    assert w['Evaluaciones']['L2'].value==0
    assert w['Evaluaciones']['L2'].data_type=='n'
    assert book(c.get(URL,{'exam':first['id']}))['Evaluaciones'].max_row==66
    assert w['Evaluaciones'].freeze_panes=='D2' and w['Evaluaciones'].auto_filter.ref=='A1:O67'
    assert 'access_token' not in str(list(w['Evaluaciones'].values))


def test_details_formula_safety_dates_and_offline_ambiguity(auth_client):
    c,_,_=auth_client();exam=create_exam(c)
    a=attempt(exam,1,7.5)
    end=datetime(2026,10,23,5,30,tzinfo=dtz.utc)
    snapshot={'title':'Examen histórico','duration_minutes':60,'questions':exam['questions']}
    RecruitmentExamAttempt.objects.filter(pk=a.pk).update(name='=HYPERLINK("https://evil.invalid")',
        phone='+525500000001',answers={'1':0,'2':'=1+1\x00'},snapshot=snapshot,grading_key={'1':0,'2':'Explicación'},
        grading={'score':7.5,'questions':{'1':{'points':10,'reason':'Correcta','review':False},'2':{'points':5,'reason':'Revisar','review':True}}},
        started_at=end-timedelta(minutes=30),submitted_at=end)
    review=attempt(exam,2,None)
    snapshot={**snapshot,'offline_import':{'source':'pdf_manual','completed_on':'2026-10-20','ambiguous_answers':{'1':['A','B']}}}
    RecruitmentExamAttempt.objects.filter(pk=review.pk).update(snapshot=snapshot,grading_status='review',
        grading={'review':True,'error':'Dos marcas'},grading_key={'1':0,'2':'Criterio'},answers={'2':'Texto'},
        started_at=end,submitted_at=end)
    w=book(c.get(URL))
    result=w['Evaluaciones'];details=w['Respuestas']
    assert result['C2'].data_type=='s' and result['C2'].value.startswith('=HYPERLINK')
    assert result['E2'].data_type=='s' and result['E2'].value=='+525500000001'
    assert result['F2'].value.date().isoformat()=='2026-10-22'
    assert result['H2'].value==datetime(2026,10,22,23,30)
    assert result['I2'].value==30 and result['L2'].value==7.5
    assert result['F3'].value.date().isoformat()=='2026-10-20'
    assert all(result[f'{col}3'].value is None for col in ['G','H','I','L'])
    assert result['K3'].value=='Revisión manual'
    assert details['H2'].value=='A. Primera' and details['I2'].value=='A. Primera'
    assert details['H3'].value=='=1+1' and details['H3'].data_type=='s'
    assert details['J3'].value==5 and details['K3'].value=='Sí'
    assert details['H4'].value=='Varias opciones: A, B' and details['J4'].value is None
    assert details['K4'].value=='Sí'
    assert not any(cell.data_type=='f' or cell.hyperlink for sheet in w for row in sheet for cell in row)
    # An explicit administrative resolution preserves the two marks without a false review flag.
    RecruitmentExamAttempt.objects.filter(pk=review.pk).update(grading_status='graded',
        grading={'score':10,'review':False,'accepted_multiple_answers':['1'],
                 'questions':{'1':{'points':10,'reason':'Doble marca aceptada: incluye la correcta.','review':False}}})
    resolved=book(c.get(URL))
    assert resolved['Evaluaciones']['L3'].value==10 and resolved['Evaluaciones']['M3'].value=='No'
    assert resolved['Respuestas']['H4'].value=='Varias opciones: A, B'
    assert resolved['Respuestas']['J4'].value==10 and resolved['Respuestas']['K4'].value=='No'
    assert resolved['Respuestas']['L4'].value=='Doble marca aceptada: incluye la correcta.'


@pytest.mark.parametrize('role,permissions',[('site_coordinator',[]),('collaborator',[]),
    ('admin',['court_communications_only']),('admin',['player_enrollments_only'])])
def test_export_denies_unrelated_accounts(auth_client,role,permissions):
    c,_,_=auth_client(role=role,section_permissions=permissions)
    assert c.get(URL).status_code in [401,403]


def test_anonymous_empty_and_bad_scope(auth_client):
    assert APIClient().get(URL).status_code in [401,403]
    c,_,_=auth_client()
    assert book(c.get(URL))['Evaluaciones'].max_row==1
    for params in [{'exam':'bad'},{'exam':'0'},{'page':'1'},{'status':'graded'}]:
        assert c.get(URL,params).status_code==400
    assert c.get(URL,{'exam':'999999'}).status_code==404
