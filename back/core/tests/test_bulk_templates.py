import io
import json
from unittest.mock import patch
import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import Workbook
from rest_framework.response import Response
from core.api.bulk_import import import_recipients, normalize_phone

def file(name, content):
    return SimpleUploadedFile(name, content)

def test_text_comma_newline_dedup_and_only_ten_digits():
    r = import_recipients(text='5512345678, 55 1234 5678\n5587654321; +525512345678,123')
    assert r['phones'] == ['5512345678', '5587654321']
    assert r['duplicates'] == 1 and len(r['invalid']) == 2

@pytest.mark.parametrize('value', ['+525512345678', '525512345678', '123', '0123456789', '=5512345678', '5.512e9', True])
def test_invalid_phone(value):
    with pytest.raises(ValueError): normalize_phone(value)

def test_csv_header_and_semicolon():
    r = import_recipients(file('numbers.csv', 'nombre;teléfono\nUno;5512345678\nDos;5587654321'.encode()))
    assert r['phones'] == ['5512345678', '5587654321']
    assert r['names'] == {'5512345678': 'Uno', '5587654321': 'Dos'}

def test_names_follow_valid_phone_after_dedup_and_no_formula():
    data = 'Nombre,Teléfono\n,5512345678\nSantiago,5512345678\nOtro,5512345678\nIncorrecto,123\n=SUM(1),5587654321'
    r = import_recipients(file('names.csv', data.encode()))
    assert r['names'] == {'5512345678': 'Santiago'}
    assert r['duplicates'] == 2

def test_excel_names_and_reversed_columns():
    w = Workbook(); w.active.append(['Teléfono', 'Nombre']); w.active.append([5512345678, '  María   López '])
    b = io.BytesIO(); w.save(b)
    assert import_recipients(file('n.xlsx', b.getvalue()))['names'] == {'5512345678': 'María López'}

def test_excel_prefills_every_matching_template_field_by_phone():
    data = 'Telefono,Nombre del destinatario,Empresa\n5512345678,Ana,Empresa Uno\n5587654321,Luis,Empresa Dos'
    parameters = json.dumps([
        {'key': 'body:1', 'label': 'Nombre del destinatario', 'contact_name': True},
        {'key': 'body:2', 'label': 'Empresa', 'contact_name': False},
    ])
    result = import_recipients(file('plantilla.csv', data.encode()), template_parameters=parameters)
    assert result['parameter_values'] == {
        '5512345678': {'body:1': 'Ana', 'body:2': 'Empresa Uno'},
        '5587654321': {'body:1': 'Luis', 'body:2': 'Empresa Dos'},
    }

def test_recruitment_columns_are_imported_by_phone():
    data = 'Telefono,Nombre,Plataforma,Tipo de vacante\n5512345678,Ana,OCC,Coach\n5587654321,Luis,Indeed,Contador'
    result = import_recipients(file('rh.csv', data.encode()))
    assert result['filters'] == {
        '5512345678': {'platform': 'OCC', 'vacancy_type': 'Coach'},
        '5587654321': {'platform': 'Indeed', 'vacancy_type': 'Contador'},
    }

def test_multiple_columns_require_selection():
    content = b'Nombre,Contacto\nUno,5512345678'
    assert import_recipients(file('n.csv', content))['needs_column']
    assert import_recipients(file('n.csv', content), column='1')['phones'] == ['5512345678']

def test_excel_numeric_cells_and_formula_not_executed():
    w = Workbook(); w.active.append(['telefono']); w.active.append([5512345678]); w.active.append(['=5511111111'])
    b = io.BytesIO(); w.save(b)
    result = import_recipients(file('n.xlsx', b.getvalue()))
    assert result['phones'] == ['5512345678'] and len(result['invalid']) == 1

def test_analyzed_workbook_uses_contactos_sheet_and_preserves_filter_metadata():
    w = Workbook()
    w.active.title = 'Filtros'
    w.active.append(['Panel visual; no contiene teléfonos'])
    contacts = w.create_sheet('Contactos')
    contacts.append(['Título'])
    contacts.append(['Descripción'])
    contacts.append([])
    contacts.append([])
    contacts.append([
        'Teléfono', 'Contacto', 'Estatus campaña', 'Relación academia',
        'Hubo interacción de ambos lados', 'Pendiente de nuestra respuesta',
        'Inició la conversación', 'Periodo de última interacción',
        'Antigüedad de la última interacción', 'Nuestro equipo respondió',
    ])
    contacts.append([
        '525512345678', 'Martha', 'Contacto permitido', 'Prospecto',
        'Sí', 'No', 'El contacto', '2026-08', '01 · Últimos 30 días', 'Sí',
    ])
    b = io.BytesIO(); w.save(b)
    result = import_recipients(file('analisis.xlsx', b.getvalue()))
    assert result['phones'] == ['5512345678']
    assert result['names'] == {'5512345678': 'Martha'}
    assert result['file_profile'] == 'academy'
    assert result['file_contacts'][0] == {
        'phone': '5512345678', 'name': 'Martha', 'campaign_status': 'Contacto permitido',
        'mutual_interaction': 'Sí', 'pending_response': 'No', 'academy_relationship': 'Prospecto',
        'league_relevance': '', 'age_bucket': '01 · Últimos 30 días', 'started_by': 'El contacto',
        'team_responded': 'Sí', 'period': '2026-08', 'year': '2026', 'month': '08',
    }

def test_league_analysis_is_identified_without_supabase_contact():
    data = (
        'Telefono,Contacto,Estatus campana,Relevancia para liga,Periodo de ultima interaccion\n'
        '5215587654321,Equipo Norte,Revisar consentimiento,Prospecto de liga,2025-11\n'
    ).encode()
    result = import_recipients(file('liga.csv', data))
    assert result['file_profile'] == 'league'
    assert result['phones'] == ['5587654321']
    assert result['file_facets']['league_relevance'] == ['Prospecto de liga']

def test_txt_and_limits_and_bad_file():
    assert import_recipients(file('n.txt', b'5512345678,5587654321'))['count'] == 2
    with pytest.raises(ValueError): import_recipients(text=','.join(['5512345678']*1001))
    with pytest.raises(ValueError): import_recipients(file('n.xlsx', b'broken'))
    with pytest.raises(ValueError): import_recipients(file('n.exe', b'a'))
    with pytest.raises(ValueError): import_recipients(file('n.csv', b'a'*(2*1024*1024+1)))

@pytest.mark.django_db
def test_veronica_only_scope_and_import(auth_client):
    client, _, _ = auth_client(role='collaborator', section_permissions=['veronica_only'])
    r = client.post('/api/veronica/bulk/import/', {'text': '5512345678,5587654321'}, format='json')
    assert r.status_code == 200 and r.data['count'] == 2
    for op in ('channels', 'catalog', 'list', 'detail', 'create', 'start', 'cancel', 'import'):
        assert client.get('/api/whatsapp-bulk/'+op+'/').status_code == 403
        assert client.post('/api/whatsapp-bulk/'+op+'/', {}, format='json').status_code == 403
    with patch('core.api.bulk.BulkView.forward', return_value=Response({})) as forward:
        client.post('/api/veronica/bulk/create/', {'channel_kind': 'academy', 'actor_id': 999, 'phones': []}, format='json')
        assert forward.call_args.args[0] == 'veronica'
        assert 'channel_kind' not in forward.call_args.kwargs['data']
        assert forward.call_args.kwargs['data']['actor_id'] != 999

@pytest.mark.django_db
def test_only_admin_and_vero_can_use_bulk(auth_client):
    for role in ('collaborator', 'site_coordinator'):
        client, _, _ = auth_client(role=role)
        assert client.post('/api/whatsapp-bulk/import/', {'text': '5512345678'}).status_code == 403
    client, _, _ = auth_client(role='admin')
    for prefix in ('/api/veronica/bulk/', '/api/whatsapp-bulk/'):
        assert client.post(prefix+'import/', {'text': '5512345678'}).status_code == 200

@pytest.mark.django_db
def test_import_enriches_missing_names_and_preserves_excel_names(auth_client):
    client, _, _ = auth_client(role='admin')
    with patch('core.api.bulk.BulkView.forward', return_value=Response({'names': {'5512345678': 'Saved', '5587654321': 'Known'}})) as lookup:
        result = client.post('/api/whatsapp-bulk/import/', {'channel': 'meta:123', 'file': file('n.csv', b'Nombre,Telefono\nExcel,5512345678\n,5587654321')}, format='multipart')
        assert result.status_code == 200
        assert result.data['names'] == {'5512345678': 'Excel', '5587654321': 'Known'}
        assert lookup.call_args.args == ('academy', 'names')
        assert lookup.call_args.kwargs['data']['channel'] == 'meta:123'

@pytest.mark.django_db
def test_veronica_import_creates_unknown_filter_categories(auth_client):
    from core.models import VeronicaFilterOption
    client, _, _ = auth_client(role='admin')
    data = b'Telefono,Plataforma,Puesto\n5512345678,LinkedIn,Recepcionista'
    result = client.post('/api/veronica/bulk/import/', {'file': file('rh.csv', data)}, format='multipart')
    assert result.status_code == 200
    assert result.data['added_filter_options'] == {'platform': ['LinkedIn'], 'vacancy_type': ['Recepcionista']}
    assert {'LinkedIn', 'Recepcionista'} <= set(VeronicaFilterOption.objects.values_list('label', flat=True))
