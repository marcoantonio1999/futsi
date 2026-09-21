from collections import defaultdict
from datetime import date, datetime, time
from io import BytesIO
import re

from django.conf import settings
from django.db import connection
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill


EXCEL_CELL_LIMIT = 32_767
HEADER_FILL = PatternFill("solid", fgColor="285943")
HEADER_FONT = Font(color="FFFFFF", bold=True)
UVM_DIRECTORY_PHONE_ID = "1078241325382505"


def _safe_text(value) -> str:
    text = ILLEGAL_CHARACTERS_RE.sub("", str(value or ""))[:EXCEL_CELL_LIMIT]
    if text.startswith(("=", "+", "-", "@")):
        return "'" + text
    return text


def _contact_name(conversation) -> str:
    context = conversation.context if isinstance(conversation.context, dict) else {}
    return (
        str(context.get("contact_name") or "").strip()
        or (
            str(conversation.booking.responsible_name or "").strip()
            if conversation.booking_id and conversation.booking
            else ""
        )
        or "Contacto de WhatsApp"
    )


def _local_excel_datetime(value):
    if not value:
        return None
    return timezone.localtime(value).replace(tzinfo=None)


def _excel_date(value):
    if isinstance(value, datetime):
        return _local_excel_datetime(value) if timezone.is_aware(value) else value
    return datetime.combine(value, time.min) if isinstance(value, date) else None


def _phone10(value) -> str:
    digits = re.sub(r"[^0-9]", "", str(value or ""))
    return digits[-10:] if len(digits) >= 10 else ""


def _contact_phone(value) -> str:
    phone = _phone10(value)
    return "+52" + phone if phone else str(value or "").strip()


def _count(value) -> int:
    try:
        return max(0, int(value or 0))
    except (TypeError, ValueError):
        return 0


def _directory_phone_id(address: str) -> str:
    if address.startswith("meta:") and address[5:].isdigit():
        return address[5:]
    configured = "".join(character for character in settings.META_WHATSAPP_DISPLAY_NUMBER if character.isdigit())
    if _phone10(address) == _phone10(configured):
        # The operational Meta id changed after the historical UVM directory was
        # imported; this is the stable id attached to that stored dataset.
        return UVM_DIRECTORY_PHONE_ID
    return ""


def directory_contacts_for_addresses(addresses, channel_details=None):
    """Read the persisted FUTSI contact directory for the selected channels."""
    channel_details = channel_details or {}
    phone_id_to_address = {
        phone_id: address
        for address in addresses
        if (phone_id := _directory_phone_id(address))
    }
    if not phone_id_to_address:
        return []
    tables = set(connection.introspection.table_names())
    if not {"whatsapp_contact_datasets", "whatsapp_contact_records"}.issubset(tables):
        return []
    placeholders = ", ".join(["%s"] * len(phone_id_to_address))
    with connection.cursor() as cursor:
        cursor.execute(
            f"""
            SELECT d.phone_number_id, r.id, r.phone, r.manual_name, r.name,
                   r.first_date, r.last_date, r.message_count, r.source_data
              FROM whatsapp_contact_records r
              JOIN whatsapp_contact_datasets d ON d.id = r.dataset_id
             WHERE d.phone_number_id IN ({placeholders})
             ORDER BY d.phone_number_id, r.ordinal, r.id
            """,
            list(phone_id_to_address),
        )
        records = cursor.fetchall()
    result = []
    for phone_id, record_id, phone, manual_name, name, first_date, last_date, message_count, source_data in records:
        address = phone_id_to_address[str(phone_id)]
        details = channel_details.get(address, {})
        source = source_data if isinstance(source_data, dict) else {}
        result.append({
            "business_address": address,
            "channel_label": details.get("label", ""),
            "site_name": details.get("site", ""),
            "directory_record_id": record_id,
            "contact_phone": _contact_phone(phone),
            "name": str(manual_name or name or "").strip(),
            "received": _count(source.get("Recibidos")),
            "sent": _count(source.get("Enviados")),
            "message_count": _count(message_count),
            "first_date": first_date,
            "last_date": last_date,
        })
    return result


def _prepare_sheet(sheet, headers, widths):
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{sheet.cell(row=1, column=len(headers)).column_letter}1"
    sheet.sheet_view.showGridLines = False
    for column, (header, width) in enumerate(zip(headers, widths), start=1):
        cell = sheet.cell(row=1, column=column, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        sheet.column_dimensions[cell.column_letter].width = width
    sheet.row_dimensions[1].height = 24


def build_whatsapp_chat_export(conversations, directory_contacts=()):
    grouped = defaultdict(lambda: {
        "label": "",
        "site": "",
        "name": "",
        "contact_phone": "",
        "conversation_ids": [],
        "messages": [],
        "directory_record": False,
        "directory_received": 0,
        "directory_sent": 0,
        "directory_first": None,
        "directory_last": None,
    })
    conversation_count = 0
    for conversation in conversations:
        conversation_count += 1
        contact_phone = _contact_phone(conversation.contact_phone)
        key = (conversation.to_address, contact_phone)
        row = grouped[key]
        row["contact_phone"] = contact_phone
        row["label"] = row["label"] or str(getattr(conversation, "channel_label", "") or "")
        row["site"] = row["site"] or str(getattr(conversation, "channel_site_name", "") or "")
        candidate_name = _contact_name(conversation)
        if not row["name"] or row["name"] == "Contacto de WhatsApp":
            row["name"] = candidate_name
        row["conversation_ids"].append(conversation.pk)
        row["messages"].extend(list(conversation.messages.all()))

    for contact in directory_contacts:
        business_address = str(contact.get("business_address") or "").strip()
        contact_phone = _contact_phone(contact.get("contact_phone"))
        if not business_address:
            continue
        # Historical imports may contain intentionally retained rows without a
        # usable phone. Keep each of them in the workbook instead of silently
        # dropping it; only records with a phone participate in deduplication.
        directory_key = contact_phone or f"__directory__:{contact.get('directory_record_id')}"
        row = grouped[(business_address, directory_key)]
        row["contact_phone"] = row["contact_phone"] or contact_phone
        row["label"] = row["label"] or str(contact.get("channel_label") or "")
        row["site"] = row["site"] or str(contact.get("site_name") or "")
        candidate_name = str(contact.get("name") or "").strip()
        if candidate_name and (not row["name"] or row["name"] == "Contacto de WhatsApp"):
            row["name"] = candidate_name
        row["directory_record"] = True
        # Repeated source records for one number must not inflate historical
        # counts. The operational conversation counts are added separately.
        row["directory_received"] = max(row["directory_received"], _count(contact.get("received")))
        row["directory_sent"] = max(row["directory_sent"], _count(contact.get("sent")))
        first_date, last_date = contact.get("first_date"), contact.get("last_date")
        if first_date and (not row["directory_first"] or first_date < row["directory_first"]):
            row["directory_first"] = first_date
        if last_date and (not row["directory_last"] or last_date > row["directory_last"]):
            row["directory_last"] = last_date

    workbook = Workbook()
    summary = workbook.active
    summary.title = "Conversaciones"
    summary_headers = [
        "Número de atención",
        "Identificador de atención",
        "Sede",
        "Número de contacto",
        "Nombre del contacto",
        "Conversaciones",
        "Mensajes recibidos",
        "Mensajes enviados",
        "Primera interacción",
        "Última interacción",
    ]
    _prepare_sheet(summary, summary_headers, [24, 29, 22, 20, 32, 15, 19, 18, 21, 21])

    detail = workbook.create_sheet("Mensajes")
    detail_headers = [
        "Número de atención",
        "Identificador de atención",
        "Número de contacto",
        "Nombre del contacto",
        "Fecha y hora",
        "Dirección",
        "Mensaje",
        "Estado",
    ]
    _prepare_sheet(detail, detail_headers, [24, 29, 20, 32, 21, 14, 78, 14])

    stored_message_rows = 0
    total_messages = 0
    ordered_groups = sorted(
        grouped.items(),
        key=lambda item: (
            item[1]["label"].casefold(),
            item[1]["name"].casefold(),
            item[0][1],
        ),
    )
    for (business_address, _group_key), row in ordered_groups:
        contact_phone = row["contact_phone"]
        messages = sorted(row["messages"], key=lambda message: (message.created_at, message.pk))
        live_received = sum(message.direction == "inbound" for message in messages)
        live_sent = sum(message.direction == "outbound" for message in messages)
        received = row["directory_received"] + live_received
        sent = row["directory_sent"] + live_sent
        stored_message_rows += len(messages)
        total_messages += received + sent
        channel_label = row["label"] or business_address
        first_values = [_excel_date(value) for value in (row["directory_first"], messages[0].created_at if messages else None) if value]
        last_values = [_excel_date(value) for value in (row["directory_last"], messages[-1].created_at if messages else None) if value]
        summary.append([
            _safe_text(channel_label),
            _safe_text(business_address),
            _safe_text(row["site"]),
            _safe_text(contact_phone),
            _safe_text(row["name"]),
            max(1 if row["directory_record"] else 0, len(row["conversation_ids"])),
            received,
            sent,
            min(first_values) if first_values else None,
            max(last_values) if last_values else None,
        ])
        for message in messages:
            revoked = str(message.body or "").strip().casefold() in {"[revoke]", "mensaje eliminado"}
            detail.append([
                _safe_text(channel_label),
                _safe_text(business_address),
                _safe_text(contact_phone),
                _safe_text(row["name"]),
                _local_excel_datetime(message.created_at),
                "Recibido" if message.direction == "inbound" else "Enviado",
                "Mensaje eliminado" if revoked else _safe_text(message.body),
                "Eliminado" if revoked else "Visible",
            ])

    for sheet in (summary, detail):
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=cell.column in (5, 7))
        for date_column in ((9, 10) if sheet is summary else (5,)):
            for cell in sheet.iter_cols(min_col=date_column, max_col=date_column, min_row=2):
                for value in cell:
                    value.number_format = "dd/mm/yyyy hh:mm"

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue(), {
        "contacts": len(grouped),
        "conversations": conversation_count,
        "messages": total_messages,
        "message_rows": stored_message_rows,
    }
