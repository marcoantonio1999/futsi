"""On-demand, private XLSX export. Never expires or grades an attempt."""
from datetime import date, datetime
from io import BytesIO
import math
from zoneinfo import ZoneInfo

from django.http import HttpResponse
from django.shortcuts import get_object_or_404
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from rest_framework.exceptions import ValidationError
from rest_framework.views import APIView

from core.api.exams import CanManageExams
from core.exam_models import RecruitmentExam, RecruitmentExamAttempt
from core.exam_profiles import filtered_attempts
from core.exam_skills import SKILLS, attempt_skills

CDMX = ZoneInfo('America/Mexico_City')
STATUSES = {'graded': 'Calificado', 'review': 'Revisión manual', 'pending': 'Pendiente de calificar',
            'processing': 'En proceso / revisar', 'error': 'Error de calificación'}


def local_clock(value):
    return value.astimezone(CDMX).replace(tzinfo=None) if value else None


def numeric(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value)


def answer_text(question, value):
    if value is None:
        return 'Sin respuesta'
    if question.get('kind') == 'choice':
        options = question.get('options', [])
        if isinstance(value, int) and not isinstance(value, bool) and 0 <= value < len(options):
            return f'{chr(65 + value)}. {options[value]}'
        return 'Respuesta inválida / revisar'
    return str(value)


def append_row(sheet, values, header=False):
    cells = []
    for value in values:
        # Applicant-controlled text must never become an Excel formula/link.
        if isinstance(value, str):
            value = ILLEGAL_CHARACTERS_RE.sub('', value)
            if len(value) > 32767:
                raise ValidationError({'detail': 'Un campo excede el tamaño de celda de Excel. No se descargó un archivo incompleto.'})
        cell = WriteOnlyCell(sheet, value=value)
        if isinstance(value, str):
            cell.data_type = 's'
        cell.font = Font(name='Calibri', size=11, color='FFFFFF' if header else '17364D', bold=header)
        cell.alignment = Alignment(vertical='top', wrap_text=True)
        if header:
            cell.fill = PatternFill('solid', fgColor='107C41')
        elif isinstance(value, datetime):
            cell.number_format = 'dd/mm/yyyy hh:mm:ss'
        elif isinstance(value, date):
            cell.number_format = 'dd/mm/yyyy'
        elif isinstance(value, float):
            cell.number_format = '0.##'
        cells.append(cell)
    sheet.append(cells)


class ExamExportView(APIView):
    permission_classes = [CanManageExams]

    def get(self, request):
        if set(request.query_params) - {'exam', 'ordering', 'top', 'q', 'status', 'min_score', 'max_score', 'skill'}:
            raise ValidationError({'detail': 'Filtro de exportación no reconocido. La descarga incluye todas las páginas.'})
        exam_id = request.query_params.get('exam')
        if exam_id is not None:
            if not exam_id.isdecimal() or int(exam_id) < 1:
                raise ValidationError({'detail': 'Examen inválido.'})
            get_object_or_404(RecruitmentExam, pk=int(exam_id))
        if request.query_params.get('status', '') not in {'', 'graded', 'pending', 'unsubmitted'} or request.query_params.get('top', '') not in {'', '10'}:
            raise ValidationError({'detail': 'Filtro de exportación inválido.'})
        if not exam_id and set(request.query_params) - {'exam'}:
            raise ValidationError({'detail': 'Selecciona un examen para exportar con filtros.'})
        if exam_id:
            rows, _ = filtered_attempts(RecruitmentExamAttempt.objects.filter(exam_id=int(exam_id)), request.query_params)
        else:
            rows = RecruitmentExamAttempt.objects.filter(submitted_at__isnull=False).order_by('exam_id', 'id')
        rows = rows.select_related('exam').defer('access_token', 'start_key', 'identity_email', 'identity_phone',
            'exam__public_token', 'exam__answer_key')
        workbook = Workbook(write_only=True)
        results = workbook.create_sheet('Evaluaciones')
        answers = workbook.create_sheet('Respuestas')
        result_headers = ['ID entrega', 'Examen / puesto', 'Participante', 'Correo', 'Teléfono', 'Fecha del examen (CDMX)',
            'Inicio (CDMX)', 'Entrega (CDMX)', 'Duración (min)', 'Tiempo de referencia (min)', 'Estado',
            'Calificación / 10', 'Requiere revisión', 'Observaciones', 'Origen']
        result_headers += [f'{label} / 10' for label, _ in SKILLS.values()]
        skill = request.query_params.get('skill', '')
        if skill:
            result_headers += ['Habilidad seleccionada', 'Calificación de habilidad / 10']
        answer_headers = ['ID entrega', 'Examen / puesto', 'Participante', 'Pregunta', 'Tipo', 'Enunciado',
            'Opciones', 'Respuesta del participante', 'Respuesta correcta / criterio', 'Puntos / 10', 'Requiere revisión', 'Comentario']
        for sheet, headers, widths in [(results, result_headers, [14,34,32,36,20,24,24,24,20,24,28,22,22,65,24] + [32] * len(SKILLS) + ([36,30] if skill else [])),
                                      (answers, answer_headers, [14,34,32,14,20,65,65,65,65,18,22,65])]:
            sheet.freeze_panes = 'D2'
            sheet.row_dimensions[1].height = 34
            for index, width in enumerate(widths, 1):
                sheet.column_dimensions[get_column_letter(index)].width = width
            append_row(sheet, headers, header=True)
        count = detail_count = 0
        for attempt in rows.iterator(chunk_size=200):
            count += 1
            if count >= 1048576:
                raise ValidationError({'detail': 'El archivo excede el límite de filas de Excel. Exporta por examen.'})
            snapshot, grading = attempt.snapshot or {}, attempt.grading or {}
            submitted = attempt.submitted_at is not None
            metadata = snapshot.get('offline_import') or {}
            offline = metadata.get('source') in {'pdf_manual', 'xlsx_results'}
            totals_only = metadata.get('source') == 'xlsx_results'
            ambiguous = metadata.get('ambiguous_answers') or {}
            accepted = set(grading.get('accepted_multiple_answers') or []) if attempt.grading_status == 'graded' else set()
            title = snapshot.get('title') or attempt.exam.title
            score = grading.get('score') if submitted and attempt.grading_status == 'graded' else None
            if not numeric(score) or not 0 <= score <= 10:
                score = None
            review = bool(grading.get('review') or (set(ambiguous) - accepted) or attempt.grading_status in {'review','error'}
                          or (attempt.grading_status == 'graded' and score is None))
            completed = (date.fromisoformat(metadata['completed_on']) if metadata.get('completed_on') else None) if offline else attempt.submitted_at.astimezone(CDMX).date() if submitted else None
            observations = (f"{metadata.get('correct_count')} aciertos de {metadata.get('question_count')}. Nota calculada con aciertos. Sin fecha, contactos, tiempos ni respuestas individuales en el archivo." if totals_only else 'Sin horas ni duración registradas' if offline else '')
            areas = attempt_skills(snapshot.get('questions'), attempt.grading_status, grading.get('questions')) if submitted else {}
            skill_scores = [areas.get(key, (None, 0))[0] for key in SKILLS]
            append_row(results, [attempt.id, title, attempt.name, attempt.email, attempt.phone, completed,
                None if offline else local_clock(attempt.started_at), None if offline else local_clock(attempt.submitted_at),
                None if offline or not submitted else round(max(0,(attempt.submitted_at-attempt.started_at).total_seconds())/60,2),
                None if totals_only else snapshot.get('duration_minutes'), STATUSES.get(attempt.grading_status, 'Revisar estado') if submitted else 'Sin entregar', score,
                'Sí' if review else 'No', grading.get('error') or observations,
                'Carga de Excel' if totals_only else 'Carga manual de PDF' if offline else 'Formulario web'] + skill_scores + ([SKILLS[skill][0], attempt.skill_score] if skill else []))
            # Unsubmitted drafts are not delivered answers and must not be exported.
            if not submitted:
                continue
            for question in snapshot.get('questions', []):
                detail_count += 1
                if detail_count >= 1048576:
                    raise ValidationError({'detail': 'El archivo excede el límite de filas de Excel. Exporta por examen.'})
                qid = question['id']
                marks = ambiguous.get(qid)
                grade = (grading.get('questions') or {}).get(qid, {})
                points = grade.get('points') if attempt.grading_status == 'graded' else None
                if not numeric(points) or not 0 <= points <= 10:
                    points = None
                append_row(answers, [attempt.id, title, attempt.name, qid,
                    'Opción múltiple' if question.get('kind') == 'choice' else 'Abierta', question.get('prompt',''),
                    '\n'.join(f'{chr(65+i)}. {option}' for i,option in enumerate(question.get('options', []))),
                    f'Varias opciones: {", ".join(marks)}' if marks else answer_text(question, attempt.answers.get(qid)),
                    answer_text(question, attempt.grading_key.get(qid)) if qid in attempt.grading_key else 'Sin clave registrada',
                    points, 'Sí' if (marks and qid not in accepted) or grade.get('review') else 'No',
                    'Confirmar selección con Vero' if marks and qid not in accepted else grade.get('reason','')])
        results.auto_filter.ref = f'A1:{get_column_letter(len(result_headers))}{max(1,count+1)}'
        answers.auto_filter.ref = f'A1:L{max(1,detail_count+1)}'
        output = BytesIO()
        workbook.save(output)
        scope = f'examen-{int(exam_id)}' if exam_id else 'todos-los-examenes'
        response = HttpResponse(output.getvalue(),content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = f'attachment; filename="{scope}-{timezone.now().astimezone(CDMX):%Y-%m-%d}.xlsx"'
        response['Cache-Control'] = 'private, no-store'
        response['X-Content-Type-Options'] = 'nosniff'
        return response
